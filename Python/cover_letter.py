import sys
import os
from docx import Document
from datetime import date

# Retrieve the path to the Python installation
python_path = sys.executable

# Prompt the user for the Word document template path
template_path = input("Enter the path of the Word document template: ").strip('\'"')

# Open the Word document template
document = Document(template_path)

# Prompt the user for recipient details
recipient_name = input("Enter the recipient's name: ")
recipient_company = input("Enter the recipient's company name: ")
recipient_address = input("Enter the recipient's address: ")
job_title = input("Enter the job title: ")

# Get today's date
today = date.today().strftime("%B %d, %Y")

# Replace placeholders in the Word document with recipient details and date
for paragraph in document.paragraphs:
    if '[RECIPIENT_NAME]' in paragraph.text:
        paragraph.text = paragraph.text.replace('[RECIPIENT_NAME]', recipient_name)
    if '[RECIPIENT_COMPANY]' in paragraph.text:
        paragraph.text = paragraph.text.replace('[RECIPIENT_COMPANY]', recipient_company)
    if '[RECIPIENT_ADDRESS]' in paragraph.text:
        paragraph.text = paragraph.text.replace('[RECIPIENT_ADDRESS]', recipient_address)
    if '[JOB_TITLE]' in paragraph.text:
        paragraph.text = paragraph.text.replace('[JOB_TITLE]', job_title)
    if '[DATE]' in paragraph.text:
        paragraph.text = paragraph.text.replace('[DATE]', today)

# Set the output path as the current directory
output_path = os.path.join(os.getcwd(), 'output.docx')

# Save the modified document with recipient details and date
document.save(output_path)
